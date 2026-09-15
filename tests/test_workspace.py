"""Production, persistence and document integration tests; no live external API calls."""
import base64
import copy
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

import engine
import production
import local_store
import file_formats
import draft_assistant

ROOT=Path(__file__).resolve().parents[1]


def sample():
    r=production.row('Печать <ПВХ>','м²','2.5','work','200.5','Внутренняя норма')
    r.update(source='Внутренний прайс цеха',price_date='2026-09-12',confirmed=True)
    return local_store.normalize_project({'project':{'title':'Проверочный заказ','client':'ООО Пример','quantity':'5','notes':'Пять табличек с печатью'},
       'rows':[r],'settings':{'profit_percent':'20','tax_mode':'none','scope_confirmed':True,'cost_basis_confirmed':True}})


class WorkspaceTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.store=local_store.LocalStore(self.tmp.name)

    def test_catalog_survives_reopening_database(self):
        r=self.store.save('catalog','Тариф',sample()['rows'][0])
        again=local_store.LocalStore(self.tmp.name).get('catalog',r['id'])
        self.assertEqual(again['data']['price'],'200.5')
        self.assertEqual(self.store.list('catalog','пвх')[0]['id'],r['id'])

    def test_versions_are_immutable_and_stale_save_fails(self):
        first=self.store.save('order','Первый',sample());changed=sample();changed['rows'][0]['price']='300'
        second=self.store.save('order','Второй',changed,first['id'],first['revision'])
        self.assertEqual(second['revision'],2)
        self.assertEqual(self.store.get('order',first['id'],1)['data']['rows'][0]['price'],'200.5')
        with self.assertRaisesRegex(ValueError,'другом окне'):
            self.store.save('order','Потерянные изменения',sample(),first['id'],1)
        self.assertEqual(len(self.store.versions(first['id'])),2)

    def test_restore_preserves_order_history_without_overwriting_existing(self):
        first=self.store.save('order','Заказ',sample());p=sample();p['rows'][0]['price']='301'
        self.store.save('order','Новая цена',p,first['id'],1)
        backup=self.store.backup();result=self.store.restore(backup)
        self.assertEqual(result['imported'],1)
        records=self.store.list('order');self.assertEqual(len(records),2)
        copied=next(x for x in records if x['id']!=first['id'])
        self.assertEqual(len(self.store.versions(copied['id'])),2)
        self.assertEqual(self.store.get('order',copied['id'],1)['data']['rows'][0]['price'],'200.5')

    def test_invalid_backup_is_atomic(self):
        self.store.save('order','Исходный',sample());backup=self.store.backup()
        backup['records'].append({'kind':'secret','name':'Плохая','data':{}})
        with self.assertRaises(ValueError):self.store.restore(backup)
        self.assertEqual(len(self.store.list('order')),1)

    def test_price_import_requires_reviewed_revision_and_is_atomic(self):
        row=sample()['rows'][0];row.update(supplier='Цех',article='A1')
        first=self.store.save('catalog',row['name'],row)
        changed={**row,'price':'800'};preview=self.store.import_preview([changed])
        self.assertEqual(preview[0]['old_price'],'200.5')
        self.store.save('catalog',row['name'],{**row,'price':'900'},first['id'],1)
        with self.assertRaises(ValueError):self.store.import_apply(preview)
        self.assertEqual(self.store.get('catalog',first['id'])['data']['price'],'900')
        self.store.import_apply(self.store.import_preview([changed]))
        self.assertEqual(self.store.get('catalog',first['id'])['data']['price'],'800')

    def test_unknown_keys_and_credentials_never_persist(self):
        p=sample();p['api_key']='SECRET';p['extensions']['api_key']='SECRET'
        record=self.store.save('order','Заказ',p)
        self.assertNotIn('SECRET',json.dumps(record))
        with self.assertRaises(ValueError):self.store.save('profile','Профиль',{'logo':'data:image/svg+xml;base64,QQ=='})

    def test_new_catalog_entry_during_import_requires_new_preview(self):
        row=sample()['rows'][0]
        preview=self.store.import_preview([row])
        self.store.save('catalog',row['name'],{**row,'price':'900'})
        with self.assertRaisesRegex(ValueError,'изменился'):self.store.import_apply(preview)
        self.assertEqual(len(self.store.list('catalog')),1)
        self.assertEqual(self.store.list('catalog')[0]['data']['price'],'900')


class ProductionTests(unittest.TestCase):
    def test_pvc_recipe_computes_sheet_area_and_cut_without_prices(self):
        p=production.recipe({'kind':'sign','quantity':'20','width_mm':'600','height_mm':'400','sheet_width_mm':'2030','sheet_height_mm':'3050'})
        self.assertEqual([r['quantity'] for r in p['rows']],['1','4.8','40'])
        self.assertTrue(all(r['price']=='' and not r['confirmed'] for r in p['rows']))
        self.assertIsNone(engine.calculate(p)['totals']['total'])

    def test_cards_two_sides_and_loss_are_explicit(self):
        p=production.recipe({'kind':'cards','quantity':'100','width_mm':'90','height_mm':'50','sheet_width_mm':'300','sheet_height_mm':'200','waste_percent':'10','sides':'2'})
        self.assertEqual(p['rows'][1]['quantity'],str(int(p['rows'][0]['quantity'])*2))
        with self.assertRaises(ValueError):production.recipe({'kind':'sign','quantity':'1','width_mm':'400','height_mm':'400','sheet_width_mm':'100','sheet_height_mm':'100'})

    def test_lighted_recipes_leave_electrical_and_contour_norms_unknown(self):
        p=production.recipe({'kind':'lightbox','quantity':'1','width_mm':'2000','height_mm':'1000'})
        self.assertEqual(p['rows'][3]['quantity'],'')
        p=production.recipe({'kind':'letters','quantity':'1'})
        self.assertEqual(p['rows'][0]['quantity'],'')

    def test_operation_setup_is_counted_once_and_money_uses_engine(self):
        r=production.operation({'name':'Печать','quantity':'20','setup_minutes':'15','minutes_per_unit':'2','hourly_rate':'2000'})
        self.assertEqual(r['minutes'],'55');self.assertEqual(r['hours'],'0.916667')
        result=engine.calculate({'rows':[r['row']]})
        self.assertEqual(result['totals']['direct'],'1833.33')

    def test_supplier_comparison_includes_pack_rounding_and_delivery(self):
        base={'supplier':'A','item':'ПВХ','unit':'лист','price':'100','delivery':'150','increment':'4','minimum':'0','date':'2026-09-12','source':'Предложение','comparable':True}
        other={**base,'supplier':'B','price':'150','increment':'','delivery':'0'}
        out=production.compare_offers({'quantity':'5','unit':'лист','offers':[base,other]})
        self.assertEqual(out['offers'][0]['total'],'950.00')
        self.assertEqual(out['best_total'],'750.00')
        self.assertTrue(out['offers'][1]['best'])

    def test_unknown_delivery_or_unconfirmed_basis_never_wins(self):
        base={'supplier':'A','item':'Материал','unit':'м²','price':'1','delivery':'','minimum':'0','date':'2026-09-12','source':'Прайс','comparable':True}
        out=production.compare_offers({'quantity':'5','unit':'м²','offers':[base,{**base,'delivery':'0','comparable':False}]})
        self.assertIsNone(out['best_total']);self.assertFalse(any(x['best'] for x in out['offers']))
        out=production.compare_offers({'quantity':'5','unit':'м²','offers':[{**base,'delivery':'0','minimum':''}]})
        self.assertIsNone(out['offers'][0]['total'])

    def test_plan_actual_preserves_unknown_and_extra_expenses(self):
        p=sample();key=p['rows'][0]['id'];p['extensions']['actuals'][key]={'amount':'600','note':'Чек'}
        self.assertIsNone(production.actuals(p)['difference'])
        p['extensions']['extra_actual']='50'
        result=production.actuals(p)
        self.assertEqual(result['known_actual'],'650.00');self.assertEqual(result['difference'],'148.75')
        p['extensions']['actuals'][key]['amount']='-10'
        with self.assertRaises(ValueError):production.actuals(p)

    def test_brief_parser_does_not_guess_dimensions_without_units(self):
        result=draft_assistant.extract({'text':'Нужно 20 табличек ПВХ 600 × 400 мм с печатью'})
        self.assertEqual(result['parameters'],{'kind':'sign','width_mm':'600','height_mm':'400','quantity':'20','installation':False,'delivery':False,'laminate':False})
        self.assertEqual(draft_assistant.extract({'text':'баннер 2x3 м тираж 4'})['parameters']['width_mm'],'2000')
        result=draft_assistant.extract({'text':'табличка 600x400'})
        self.assertEqual(result['parameters']['width_mm'],'')
        self.assertTrue(draft_assistant.extract({'text':'Баннер с монтажом, без доставки'})['parameters']['installation'])
        self.assertFalse(draft_assistant.extract({'text':'Баннер с доставкой. Доставка не нужна.'})['parameters']['delivery'])

    def test_ollama_adapter_rejects_cloud_and_validates_model_output(self):
        local={'name':'local:small','size':2000000,'details':{'format':'gguf'}}
        with patch.object(draft_assistant,'ollama',return_value={'models':[local,{'name':'foo-cloud','size':100,'details':{}}]}):
            self.assertEqual(draft_assistant.models()['models'],['local:small'])
        responses=[{'models':[local]},{'message':{'content':json.dumps({'kind':'sign','width_mm':'600','height_mm':'400','quantity':'20','price':'999'})}}]
        with patch.object(draft_assistant,'ollama',side_effect=responses) as mock:
            r=draft_assistant.extract({'text':'Задание','use_ai':True,'model':'local:small'})
            self.assertNotIn('price',r['parameters']);self.assertEqual(mock.call_args.args[1]['stream'],False)


class FileFormatTests(unittest.TestCase):
    def test_xlsx_is_readable_and_numeric_names_are_not_converted(self):
        from openpyxl import load_workbook
        p=sample();p['rows'][0]['name']='00123'
        raw=file_formats.export_xlsx(p,{'company':'=HYPERLINK("bad")'})
        book=load_workbook(BytesIO(raw),data_only=False);sheet=book.active
        values=[c for row in sheet for c in row if c.value is not None]
        self.assertTrue(any(c.value=='00123' and c.data_type=='s' for c in values))
        self.assertTrue(any(c.value=='=HYPERLINK("bad")' and c.data_type=='s' for c in values))
        self.assertFalse(any(c.data_type=='f' for c in values))
        parsed=file_formats.xlsx_read({'content':base64.b64encode(raw).decode()})
        self.assertTrue(any('00123' in r['values'] for r in parsed['rows']))

    def test_xlsx_mapping_preserves_sources_and_blocks_formula_prices(self):
        body={'rows':[{'number':'1','values':['Название','Ед.','Цена'],'formulas':[]},
                      {'number':'2','values':['ПВХ','лист','1 200,50'],'formulas':[]}],
              'mapping':{'name':0,'unit':1,'price':2},'source':'price.xlsx / Лист1','date':'2026-09-12'}
        r=file_formats.map_prices(body);self.assertEqual(r['rows'][0]['price'],'1200.5');self.assertFalse(r['errors'])
        self.assertIn('строка 2',r['rows'][0]['source'])
        body['rows'][1]['formulas']=[2];r=file_formats.map_prices(body)
        self.assertTrue(r['errors']);self.assertFalse(r['rows'])

    def test_invalid_xlsx_and_dtd_rejected(self):
        with self.assertRaises(ValueError):file_formats.xlsx_read({'content':base64.b64encode(b'not xlsx').decode()})
        raw=file_formats.export_xlsx(sample());dest=BytesIO()
        with zipfile.ZipFile(BytesIO(raw)) as src,zipfile.ZipFile(dest,'w') as z:
            for name in src.namelist():z.writestr(name,b'<!DOCTYPE unsafe>'+src.read(name) if name=='xl/workbook.xml' else src.read(name))
        with self.assertRaisesRegex(ValueError,'DTD'):file_formats.xlsx_read({'content':base64.b64encode(dest.getvalue()).decode()})

    def test_pdf_cyrillic_and_client_privacy(self):
        from pypdf import PdfReader
        p=sample();profile={'company':'ООО Проверка','details':'Реквизиты предприятия','terms':'Согласовать срок изготовления'}
        internal=file_formats.export_pdf(p,profile,False,ROOT)
        client=file_formats.export_pdf(p,profile,True,ROOT)
        text=''.join(x.extract_text() for x in PdfReader(BytesIO(internal)).pages)
        public=''.join(x.extract_text() for x in PdfReader(BytesIO(client)).pages)
        self.assertIn('Внутренний прайс цеха',text);self.assertIn('ООО Проверка',public)
        self.assertNotIn('Внутренний прайс цеха',public);self.assertNotIn('Расчётная прибыль',public)
        self.assertIn('601,50',public)

    def test_logo_embeds_in_pdf_and_xlsx(self):
        from PIL import Image
        from openpyxl import load_workbook
        buf=BytesIO();Image.new('RGB',(80,30),'navy').save(buf,format='PNG')
        profile={'company':'Компания','logo':'data:image/png;base64,'+base64.b64encode(buf.getvalue()).decode()}
        raw=file_formats.export_xlsx(sample(),profile,True)
        with zipfile.ZipFile(BytesIO(raw)) as z:self.assertIn('xl/media/logo.png',z.namelist())
        book=load_workbook(BytesIO(raw));self.assertEqual(len(book.active._images),1)
        self.assertTrue(file_formats.export_pdf(sample(),profile,True,ROOT).startswith(b'%PDF'))


if __name__=='__main__':unittest.main()

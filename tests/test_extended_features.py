import base64
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock,patch
import uuid
import zipfile
import xml.etree.ElementTree as ET
import advanced_documents
import catalog_match
import engine
import full_backup
import local_store
import secret_store
from stock import Stock

class StockTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup);self.store=local_store.LocalStore(self.temp.name);self.stock=Stock(self.store)
        self.item=self.store.save('catalog','ПВХ',{'name':'ПВХ','unit':'лист','price':'100','quantity':'1','category':'material'})
        self.order=self.store.save('order','Заказ',{'rows':[]})

    def move(self,kind,quantity,**extra):
        return self.stock.move({'request_id':uuid.uuid4().hex,'kind':kind,'quantity':quantity,'item_id':self.item['id'],'order_id':self.order['id'],'note':'Учебная операция',**extra})

    def test_reserve_consume_release_and_replay(self):
        self.move('receipt','10');request=uuid.uuid4().hex
        self.move('reserve','6',request_id=request)
        self.assertTrue(self.move('reserve','6',request_id=request)['replayed'])
        with self.assertRaises(ValueError):self.move('reserve','7',request_id=request)
        with self.assertRaises(ValueError):self.move('reserve','5')
        with self.assertRaises(ValueError):self.move('writeoff','5')
        self.move('consume','2');self.move('release','4')
        result=self.stock.status();self.assertEqual(result['items'][0]['on_hand'],'8');self.assertEqual(result['items'][0]['reserved'],'0')
        self.assertEqual(len(result['history']),4)
        self.assertEqual(self.store.get('catalog',self.item['id'])['data']['price'],'100')

    def test_concurrent_reservation_has_one_winner(self):
        self.move('receipt','10')
        def reserve():
            try:self.move('reserve','7');return True
            except ValueError:return False
        with ThreadPoolExecutor(max_workers=2) as pool:results=list(pool.map(lambda _:reserve(),range(2)))
        self.assertEqual(sum(results),1);self.assertEqual(self.stock.status()['items'][0]['available'],'3')

    def test_units_and_deletion_guards(self):
        self.move('receipt','5');self.move('reserve','2')
        with self.assertRaises(ValueError):self.stock.assert_delete_allowed('order',self.order['id'])
        with self.assertRaises(ValueError):self.stock.assert_delete_allowed('catalog',self.item['id'])
        self.store.save('catalog','ПВХ',{**self.item['data'],'unit':'м²'},self.item['id'],self.item['revision'])
        with self.assertRaisesRegex(ValueError,'Единица'):self.move('receipt','1')

    def test_full_backup_and_fuzzy_suggestion(self):
        self.move('receipt','5');directory=self.store.directory/'price-documents';directory.mkdir();(directory/'original.bin').write_bytes(b'original price')
        data=full_backup.create(self.store,self.stock)
        with zipfile.ZipFile(BytesIO(data)) as z:
            self.assertEqual(z.read('price-documents/original.bin'),b'original price')
            self.assertIn('stock.sqlite',z.namelist());self.assertIn('workspace.sqlite',z.namelist())
            self.assertNotIn('credentials.json',z.namelist())
        results=catalog_match.suggest(self.store,{'query':'ПВХ белый','unit':'лист'})
        self.assertEqual(results['items'][0]['id'],self.item['id'])
        self.assertEqual(catalog_match.suggest(self.store,{'query':'ПВХ белый','unit':'м²'})['items'],[])

class SecretTests(unittest.TestCase):
    def test_saved_key_is_only_used_by_explicit_request(self):
        vault=Mock();vault.get_password.return_value='tvly-private-test-key'
        with patch('secret_store.backend',return_value=vault):
            self.assertIsNone(secret_store.effective({},'tavily'));vault.get_password.assert_not_called()
            self.assertEqual(secret_store.effective({'api_key':'inline'},'tavily'),'inline')
            self.assertEqual(secret_store.effective({'use_saved_key':True},'tavily'),'tvly-private-test-key')
            status=secret_store.status();self.assertNotIn('private',json.dumps(status))
            self.assertEqual(secret_store.save('tavily','tvly-private-test-key'),{'saved':True,'service':'tavily'})

    def test_plaintext_backend_is_rejected(self):
        import keyring
        with patch('keyring.get_keyring',return_value=Mock()):
            with self.assertRaisesRegex(ValueError,'текстовые'):secret_store.backend()

class DocumentTests(unittest.TestCase):
    def demo(self,**settings):
        return {'project':{'title':'Учебная табличка <А>','client':'ООО «Пример»','quantity':'10','notes':'ПВХ, печать и резка. Учебные числа.'},'rows':[
            {'id':'one','name':'Секретная закупочная расценка','unit':'лист','quantity':'2.1','increment':'1','minimum_charge':'0','price':'100.005','source':'TEST','price_date':'2026-09-01','confirmed':True}],
            'settings':{'overhead_percent':'10','reserve_percent':'5','profit_percent':'20','discount_percent':'3','tax_mode':'add','tax_percent':'20',**settings}}

    def test_word_customer_only_and_safe_template(self):
        data=advanced_documents.export_docx(self.demo(),{'company':'Учебная мастерская'})
        with zipfile.ZipFile(BytesIO(data)) as z:
            text=''.join(ET.fromstring(z.read('word/document.xml')).itertext())
        self.assertIn('Учебная табличка <А>',text);self.assertNotIn('Секретная закупочная',text);self.assertNotIn('{{',text)
        template=advanced_documents.default_template()
        source=BytesIO()
        with zipfile.ZipFile(BytesIO(template)) as original,zipfile.ZipFile(source,'w') as modified:
            for name in original.namelist():
                data=original.read(name)
                if name=='word/document.xml':data=data.replace(b'{{ heading }}',b'{{ heading.__class__ }}')
                modified.writestr(name,data)
        with self.assertRaises(ValueError):advanced_documents.export_docx(self.demo(),template=source.getvalue())

    def test_excel_cached_values_rounding_and_blank_price(self):
        import openpyxl
        for project in (self.demo(),self.demo(pricing_mode='margin'),self.demo(tax_mode='unknown')):
            raw=advanced_documents.export_editable_xlsx(project)
            cached=openpyxl.load_workbook(BytesIO(raw),data_only=True);formula=openpyxl.load_workbook(BytesIO(raw),data_only=False)
            expected=engine.calculate(project)
            for cell,key in [('B16','direct'),('B17','overhead'),('B18','reserve'),('B19','cost'),('B20','sale_before_discount'),('B21','discount'),('B22','sale_net'),('B23','tax'),('B24','total'),('B25','profit'),('B26','unit_total')]:
                actual=cached['Параметры'][cell].value;value=expected['totals'][key]
                if value is None:self.assertIn(actual,(None,''))
                else:self.assertAlmostEqual(actual,float(value),places=2)
            self.assertTrue(formula['Строки']['G2'].value.startswith('=IF'))
            self.assertAlmostEqual(cached['Строки']['H2'].value,300.02)
        project=self.demo();project['rows'][0]['price']=''
        data=advanced_documents.export_editable_xlsx(project);book=openpyxl.load_workbook(BytesIO(data),data_only=True)
        self.assertIsNone(book['Параметры']['B24'].value);self.assertIsNone(book['Строки']['H2'].value)

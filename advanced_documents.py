"""Customer Word templates and editable Excel calculation exports."""
from decimal import Decimal
from io import BytesIO
import re
import zipfile
import xml.etree.ElementTree as ET
import engine
import file_formats
from local_store import normalize_project


def default_template():
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
    document=Document();section=document.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.left_margin=section.right_margin=Inches(.85);section.top_margin=section.bottom_margin=Inches(.7)
    for name in ('Normal','Title','Heading 1','Heading 2'):
        style=document.styles[name];style.font.name='DejaVu Sans';style.font.color.rgb=RGBColor(0,0,0)
    for style in document.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    document.styles['Normal'].font.size=Pt(10)
    document.styles['Normal'].paragraph_format.space_after=Pt(6)
    document.add_paragraph('{{ heading }}','Title')
    document.add_paragraph('{{ title }}','Heading 1')
    for value in ('{{ company }}','{{ details }}','{{ contacts }}','{{ client }}','{{ date }}','{{ status }}'):
        document.add_paragraph(value)
    table=document.add_table(rows=1,cols=4);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    widths=[3.0,1.0,1.35,1.45]
    for column,width in zip(table.columns,widths):column.width=Inches(width)
    for cell,label,width in zip(table.rows[0].cells,['Изделие или заказ','Кол-во','Средняя цена ₽','Сумма ₽'],widths):
        cell.text=label;cell.width=Inches(width)
        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6F1');cell._tc.get_or_add_tcPr().append(shade)
        for run in cell.paragraphs[0].runs:run.bold=True
    repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
    table.add_row().cells[0].text='{%tr for r in rows %}'
    for i,cell in enumerate(table.add_row().cells):cell.text='{{ r['+str(i)+'] }}'
    table.add_row().cells[0].text='{%tr endfor %}'
    for table_row in table.rows:
        for cell,width in zip(table_row.cells,widths):
            cell.width=Inches(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for side in ('top','left','bottom','right'):
                element=OxmlElement('w:'+side);element.set(qn('w:val'),'single');element.set(qn('w:sz'),'4');element.set(qn('w:color'),'D9D9D9');borders.append(element)
            properties.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ('top','left','bottom','right'):
                node=OxmlElement('w:'+side);node.set(qn('w:w'),'90');node.set(qn('w:type'),'dxa');margins.append(node)
            properties.append(margins)
    document.add_paragraph('')
    document.add_paragraph('Итого по заказу: {{ total }} ₽')
    document.add_paragraph('Начисленный НДС: {{ tax }} ₽')
    document.add_paragraph('Состав и условия','Heading 2')
    document.add_paragraph('{%p for note in notes %}')
    document.add_paragraph('{{ note }}')
    document.add_paragraph('{%p endfor %}')
    output=BytesIO();document.save(output);return output.getvalue()


def validate_template(raw):
    from jinja2 import nodes
    from jinja2.sandbox import SandboxedEnvironment
    if len(raw)>2*1024*1024:raise ValueError('Шаблон Word должен быть до 2 МБ.')
    try:
        with zipfile.ZipFile(BytesIO(raw)) as archive:
            entries=archive.infolist()
            if len(entries)>200 or sum(item.file_size for item in entries)>8*1024*1024:raise ValueError('Шаблон Word слишком большой.')
            if any('vbaProject' in item.filename or '/embeddings/' in item.filename for item in entries):raise ValueError('Шаблоны с макросами и вложенными объектами не поддерживаются.')
            for item in entries:
                if not item.filename.endswith('.xml'):continue
                content=archive.read(item).decode('utf-8')
                if '<!DOCTYPE' in content or '<!ENTITY' in content:raise ValueError('DTD/ENTITY в шаблоне не поддерживаются.')
                # Join text runs to detect expressions Word has split during edits.
                text=''.join(ET.fromstring(content).itertext())
                for candidate in (content,text):
                    candidate=re.sub(r'{%\s*(?:tr|p|tc|r)\s+', '{% ',candidate)
                    tree=SandboxedEnvironment().parse(candidate)
                    allowed=(nodes.Template,nodes.Output,nodes.TemplateData,nodes.Name,nodes.Getitem,nodes.Const,nodes.For)
                    for node in tree.find_all(nodes.Node):
                        if not isinstance(node,allowed):raise ValueError('Шаблон допускает только поля и простые циклы rows/notes.')
                        if isinstance(node,nodes.Name) and node.name not in {'heading','title','company','details','contacts','client','date','status','rows','r','notes','note','total','tax'}:
                            raise ValueError('Неизвестное поле шаблона: '+node.name)
                        if isinstance(node,nodes.For):
                            if not isinstance(node.iter,nodes.Name) or node.iter.name not in ('rows','notes') or any(isinstance(child,nodes.For) for child in node.find_all(nodes.For)):
                                raise ValueError('Вложенные и произвольные циклы в шаблонах не поддерживаются.')
            if 'word/document.xml' not in archive.namelist():raise ValueError('Не найден документ Word.')
    except (zipfile.BadZipFile,UnicodeError,ET.ParseError):raise ValueError('Не удалось прочитать шаблон DOCX.') from None


def export_docx(project,profile=None,template=None):
    try:
        from docxtpl import DocxTemplate
        from jinja2 import StrictUndefined
        from jinja2.sandbox import SandboxedEnvironment
    except ImportError:raise ValueError('Для Word установите дополнения SETUP_FEATURES.') from None
    profile,lines,columns,rows,footer,notes=file_formats.report_content(project,profile,True)
    data=template or default_template();validate_template(data)
    rendered=DocxTemplate(BytesIO(data))
    fmt=lambda value:'Не определено' if value is None or value=='' else str(value).replace('.',',') if re.fullmatch(r'[0-9]+(?:\.[0-9]+)?',str(value)) else str(value)
    context={'heading':lines[3],'title':lines[4],'company':profile['company'],'details':profile['details'],'contacts':profile['contacts'],
             'client':lines[6],'date':lines[7],'status':lines[5],'rows':[[fmt(v) for v in r] for r in rows],
             'notes':lines[8:]+notes,'total':fmt(footer[0][1]),'tax':fmt(footer[1][1])}
    try:rendered.render(context,jinja_env=SandboxedEnvironment(undefined=StrictUndefined,autoescape=True),autoescape=True)
    except Exception:raise ValueError('Не удалось заполнить шаблон Word. Используйте поля из стандартного шаблона.') from None
    if template is None:
        from docx.oxml.ns import qn
        for paragraph in list(rendered.docx.paragraphs):
            previous=paragraph._p.getprevious()
            if not paragraph.text.strip() and previous is not None and previous.tag!=qn('w:tbl'):
                paragraph._p.getparent().remove(paragraph._p)
    output=BytesIO();rendered.save(output)
    return output.getvalue()


def export_editable_xlsx(project):
    try:import xlsxwriter
    except ImportError:raise ValueError('Для Excel с формулами установите дополнения SETUP_FEATURES.') from None
    normalized=normalize_project(project);result=engine.calculate(normalized);settings=normalized['settings'];rows=result['rows']
    # Excel's numerical precision must not silently replace the exact decimal engine.
    for value in result['totals'].values():
        if value is not None and len(Decimal(value).normalize().as_tuple().digits)>15:
            raise ValueError('Итоги превышают точность Excel. Используйте расчёт программы и PDF.')
    for item in rows:
        for key in ('quantity','price','increment','minimum_charge','billed_quantity','amount'):
            if item.get(key) and len(Decimal(item[key]).normalize().as_tuple().digits)>15:
                raise ValueError('Числа выходят за точность Excel. Используйте расчёт программы и PDF/обычный экспорт.')
    output=BytesIO();book=xlsxwriter.Workbook(output,{'in_memory':True,'strings_to_formulas':False,'strings_to_urls':False})
    book.set_calc_mode('auto')
    title=book.add_format({'bold':True,'font_size':16,'font_color':'#000000'})
    heading=book.add_format({'bold':True,'bg_color':'#DCE6F1','text_wrap':True,'border':1,'border_color':'#D9D9D9','valign':'vcenter'})
    text=book.add_format({'text_wrap':True,'valign':'vcenter'})
    input_number=book.add_format({'num_format':'General','valign':'vcenter','font_color':'#0000FF','bg_color':'#FFF2CC'})
    calculated=book.add_format({'num_format':'General','valign':'vcenter','font_color':'#000000'})
    money=book.add_format({'num_format':'#,##0.00','valign':'vcenter','font_color':'#000000'})
    percent=book.add_format({'num_format':'0.00%','valign':'vcenter','font_color':'#0000FF','bg_color':'#FFF2CC'})
    params=book.add_worksheet('Параметры');calc=book.add_worksheet('Строки')
    params.hide_gridlines(2);calc.hide_gridlines(2);params.set_column('A:A',43);params.set_column('B:B',24);params.set_column('C:C',58)
    params.set_row(0,28);params.write('A1','Калькуляция заказа',title);params.write('A2',normalized['project']['title'],text)
    params.write('A3','Синие числа и жёлтые ячейки можно менять. Изменения Excel не возвращаются в программу.',text);params.set_row(2,42)
    input_rows=[('Тираж',normalized['project']['quantity']),('Накладные',settings['overhead_percent']),('Резерв',settings['reserve_percent']),('Прибыль наценка или маржа',settings['profit_percent']),('Скидка',settings['discount_percent']),('НДС ставка',settings['tax_percent']),('Ценообразование',settings['pricing_mode']),('Налоговый режим',settings['tax_mode'])]
    for i,(label,value) in enumerate(input_rows,5):
        params.write(i,0,label,text)
        if i in (11,12):params.write(i,1,value,input_number)
        elif value!='':params.write_number(i,1,float(value)/(100 if 6<=i<=10 else 1),percent if 6<=i<=10 else input_number)
        else:params.write_blank(i,1,None,percent if 6<=i<=10 else input_number)
    params.data_validation('B12',{'validate':'list','source':['markup','margin']});params.data_validation('B13',{'validate':'list','source':['none','add','unknown']})
    params.data_validation('B6',{'validate':'decimal','criteria':'>','value':0})
    params.data_validation('B7:B8',{'validate':'decimal','criteria':'between','minimum':0,'maximum':10})
    params.data_validation('B10:B11',{'validate':'decimal','criteria':'between','minimum':0,'maximum':1})
    params.write('C12','markup — наценка; margin — маржа',text);params.write('C13','none — без начисления; add — начислить; unknown — не выбран',text)
    headers=['Позиция','Единица','Количество','Цена ₽','Кратность','Минимум ₽','К оплате','Сумма ₽','Источник и дата']
    calc.write_row(0,0,headers,heading);calc.set_row(0,32);calc.freeze_panes(1,2)
    calc.set_column('A:A',33);calc.set_column('B:B',12);calc.set_column('C:H',16);calc.set_column('I:I',48)
    def cached(value):return '' if value is None else float(value)
    for i,item in enumerate(rows,1):
        r=i+1;calc.write(i,0,item['name'],text);calc.write(i,1,item['unit'],text)
        for column,key in enumerate(('quantity','price','increment','minimum_charge'),2):
            if item[key]!='':calc.write_number(i,column,float(item[key]),input_number)
            else:calc.write_blank(i,column,None,input_number)
        calc.write_formula(i,6,f'=IF(OR(NOT(ISNUMBER(C{r})),C{r}<=0),"",IF(ISBLANK(E{r}),C{r},IF(AND(ISNUMBER(E{r}),E{r}>0),CEILING(C{r},E{r}),"")))',calculated,cached(item['billed_quantity']))
        calc.write_formula(i,7,f'=IF(OR(NOT(ISNUMBER(G{r})),NOT(ISNUMBER(D{r})),NOT(ISNUMBER(F{r})),D{r}<0,F{r}<0),"",ROUND(MAX(G{r}*D{r},F{r}),2))',money,cached(item['amount']))
        calc.write(i,8,'\n'.join(filter(None,(item['source'],item['price_date'],item['note']))),text);calc.set_row(i,42)
    n=len(rows);last=max(2,n+1);known=sum(r['amount'] is not None for r in rows)==n and n>0
    formulas=[('direct','Прямые затраты',f'=ROUND(SUM(Строки!H2:H{last}),2)'),
      ('overhead','Накладные', '=IF(B16="","",ROUND(B16*B7,2))'),
      ('reserve','Резерв','=IF(OR(B17="",B8=""),"",ROUND((B16+B17)*B8,2))'),
      ('cost','Себестоимость','=IF(B18="","",ROUND(SUM(B16:B18),2))'),
      ('sale_before_discount','Цена до скидки','=IF(OR(B19="",B9=""),"",IF(B12="markup",ROUND(B19*(1+B9),2),IF(AND(B12="margin",B9<1),ROUND(B19/(1-B9),2),"")))'),
      ('discount','Скидка','=IF(OR(B20="",B10=""),"",ROUND(B20*B10,2))'),
      ('sale_net','Цена до НДС','=IF(B21="","",ROUND(B20-B21,2))'),
      ('tax','Начисленный НДС','=IF(B22="","",IF(B13="none",0,IF(AND(B13="add",ISNUMBER(B11)),ROUND(B22*B11,2),"")))'),
      ('total','Цена заказа','=IF(B23="","",ROUND(B22+B23,2))'),
      ('profit','Расчётная прибыль','=IF(B22="","",ROUND(B22-B19,2))'),
      ('unit_total','Цена за изделие','=IF(OR(B24="",B6="",B6<=0),"",ROUND(B24/B6,2))')]
    # Owning amount completeness is derived from every row, including blank vs zero.
    formulas[1]=(formulas[1][0],formulas[1][1],f'=IF(OR({n}=0,B7="",COUNT(Строки!H2:H{last})<>{n}),"",ROUND(B16*B7,2))')
    for i,(key,label,formula) in enumerate(formulas,15):
        params.write(i,0,label,heading if key=='total' else text);params.write_formula(i,1,re.sub(r'(B[0-9]+)=""',r'NOT(ISNUMBER(\1))',formula).replace('Строки!',"'Строки'!"),money,cached(result['totals'][key]))
    params.write('A29','Статус на момент выгрузки',heading);params.write('B29',result['status'],text);params.set_row(28,34)
    params.write('A30','После изменения чисел проверьте состав заказа, цены и налоги заново. Статус выгрузки не подтверждает новые данные.',text);params.set_row(29,48)
    params.write('A32','Пустая цена или неизвестный налог оставляет итог пустым. Известная часть прямых расходов показывается отдельно.',text);params.set_row(31,48)
    params.set_landscape();params.set_paper(9);params.fit_to_pages(1,1);params.print_area(0,0,31,2)
    calc.set_landscape();calc.set_paper(9);calc.fit_to_pages(1,0);calc.repeat_rows(0);calc.print_area(0,0,max(1,n),8)
    book.close();return output.getvalue()
